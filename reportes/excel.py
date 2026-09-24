"""
reportes.excel
==============

Escritura del libro Excel. Solo formato: los numeros vienen de reportes.datos.

Contrato de estructura (para automatizar la lectura despues)
------------------------------------------------------------
- Cada hoja de datos tiene UNA tabla de Excel con nombre (ListObject), p.ej.
  ``tbl_stock_clase``. Leer siempre por nombre de tabla, no por rango.
- Las tablas empiezan en la fila 6 (encabezado) de cada hoja; filas 1-4 son
  titulo, fecha de corte, descripcion y unidades.
- Lo que se ve es lo que se guarda: un 7,2% se guarda como 7,2 (columna
  "(%)"), los montos en millones de USD se guardan en millones. Ningun formato
  de celda escala el valor.
- Los totales van FUERA de la tabla, separados por una fila en blanco, para
  que sumar una columna de la tabla no cuente dos veces.
- La hoja Diccionario describe cada columna de cada tabla; se genera de las
  mismas especificaciones que arman las hojas, asi que no puede desalinearse.
"""
from __future__ import annotations

import datetime as _dt
import math
from dataclasses import dataclass
from pathlib import Path
from typing import Callable

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

# ---------------------------------------------------------------------------
#  estilo
# ---------------------------------------------------------------------------

AZUL = "1F3864"
F_TITULO = Font(name="Calibri", size=14, bold=True, color=AZUL)
F_CORTE = Font(name="Calibri", size=12, bold=True, color="000000")
F_TEXTO = Font(name="Calibri", size=10, color="404040")
F_HEADER = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
F_TOTAL = Font(name="Calibri", size=10, bold=True)
RELLENO_HEADER = PatternFill("solid", fgColor=AZUL)
RELLENO_CORTE = PatternFill("solid", fgColor="FFF2CC")
RELLENO_TOTAL = PatternFill("solid", fgColor="D9E1F2")
BORDE = Border(bottom=Side(style="thin", color="808080"))
FILA_HEADER = 6

#: Formato de celda y unidad (para el diccionario) por tipo de columna.
FORMATOS = {
    "texto": ("@", "texto"),
    "texto_largo": ("@", "texto"),
    "id": ("0", "numero sin separador"),
    "entero": ("#,##0", "cantidad"),
    "mm": ("#,##0.00", "MM USD"),
    "pct": ("0.00", "% (7,2 = 7,2%)"),
    "tasa": ("0.000", "% segun contrato"),
    "pb": ("#,##0.0", "puntos base"),
    "num0": ("#,##0", "numero"),
    "num2": ("#,##0.00", "numero"),
    "anios": ("0.00", "anios"),
    "fecha": ("dd-mm-yyyy", "fecha"),
}


@dataclass(frozen=True)
class Col:
    """Una columna de una tabla: encabezado, de donde sale, formato y definicion."""
    header: str
    fuente: str | Callable[[pd.DataFrame], pd.Series]
    fmt: str
    descripcion: str
    total: bool = False            # se suma en la fila de total fuera de la tabla


@dataclass
class Hoja:
    nombre: str                    # nombre de la pestana (<= 31 caracteres)
    tabla: str                     # nombre de la tabla de Excel
    titulo: str
    descripcion: str
    cols: list[Col]
    df: pd.DataFrame
    congelar_cols: int = 2         # columnas fijas a la izquierda
    total_etiqueta: str | None = None


def _valor(v):
    """Celda de Excel: NaN y NaT a vacio, fechas a date, numpy a python."""
    if v is None or v is pd.NA or v is pd.NaT:
        # pd.NA llega desde columnas enteras con nulos (p.ej. un RUT que solo
        # existe para los papeles nacionales).
        return None
    if isinstance(v, float) and math.isnan(v):
        return None
    if isinstance(v, (np.floating,)):
        return None if np.isnan(v) else float(v)
    if isinstance(v, (np.integer,)):
        return int(v)
    if isinstance(v, pd.Timestamp):
        return None if pd.isna(v) else v.to_pydatetime().date()
    if isinstance(v, _dt.datetime):
        return v.date()
    if v is pd.NaT:
        return None
    return v


def _columna(df: pd.DataFrame, c: Col) -> pd.Series:
    if callable(c.fuente):
        return c.fuente(df)
    if c.fuente not in df.columns:
        raise KeyError(f"La columna fuente '{c.fuente}' no existe para '{c.header}'")
    return df[c.fuente]


class Libro:
    """Arma el libro hoja por hoja y lleva el registro para el diccionario."""

    def __init__(self, corte: _dt.date, unidades: str):
        self.wb = Workbook()
        self.wb.remove(self.wb.active)
        self.corte = corte
        self.unidades = unidades
        self.registro: list[tuple[str, str, Col]] = []   # (hoja, tabla, col)

    # -- encabezado comun ------------------------------------------------------

    def _encabezado(self, ws, titulo: str, descripcion: str, unidades: str | None = None):
        ws["A1"] = titulo
        ws["A1"].font = F_TITULO
        ws["A2"] = f"Datos al: {self.corte:%d-%m-%Y}"
        ws["A2"].font = F_CORTE
        ws["A2"].fill = RELLENO_CORTE
        ws["A3"] = descripcion
        ws["A3"].font = F_TEXTO
        ws["A4"] = unidades or self.unidades
        ws["A4"].font = F_TEXTO
        ws.sheet_view.showGridLines = False

    # -- tabla -------------------------------------------------------------------

    def tabla(self, ws, df: pd.DataFrame, cols: list[Col], nombre_tabla: str,
              fila: int = FILA_HEADER, col0: int = 1, total_etiqueta: str | None = None,
              hoja: str | None = None) -> int:
        """Escribe una tabla de Excel con nombre. Devuelve la ultima fila usada."""
        headers = [c.header for c in cols]
        assert len(set(headers)) == len(headers), f"encabezados repetidos en {nombre_tabla}"
        datos = [_columna(df, c).tolist() for c in cols] if len(df) else [[] for _ in cols]
        n = len(df)

        for j, h in enumerate(headers):
            cell = ws.cell(row=fila, column=col0 + j, value=h)
            cell.font = F_HEADER
            cell.fill = RELLENO_HEADER
            cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.row_dimensions[fila].height = 42

        for j, c in enumerate(cols):
            fmt = FORMATOS[c.fmt][0]
            largo = c.fmt == "texto_largo"
            for i in range(n):
                cell = ws.cell(row=fila + 1 + i, column=col0 + j, value=_valor(datos[j][i]))
                cell.number_format = fmt
                if largo:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")

        # Una tabla de Excel necesita al menos una fila de datos.
        if n == 0:
            ws.cell(row=fila + 1, column=col0, value=None)
        ultima = fila + max(n, 1)
        ref = f"{get_column_letter(col0)}{fila}:{get_column_letter(col0 + len(cols) - 1)}{ultima}"
        t = Table(displayName=nombre_tabla, ref=ref)
        t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(t)

        # Anchos: encabezado y una muestra del contenido, con tope.
        for j, c in enumerate(cols):
            muestra = [str(_valor(v)) for v in datos[j][:300] if _valor(v) is not None]
            ancho = max([len(c.header) * 0.55 + 4] + [min(len(s), 60) + 2 for s in muestra] + [9])
            if c.fmt == "texto_largo":
                ancho = min(max(ancho, 40), 80)
            elif c.fmt in ("mm", "pct", "tasa", "pb", "num0", "num2", "anios", "entero"):
                ancho = max(min(ancho, 16), 12)
            ws.column_dimensions[get_column_letter(col0 + j)].width = min(ancho, 60)

        if total_etiqueta and n:
            ft = ultima + 2
            ws.cell(row=ft, column=col0, value=total_etiqueta).font = F_TOTAL
            for j, c in enumerate(cols):
                if c.total:
                    s = pd.to_numeric(pd.Series(datos[j]), errors="coerce").sum()
                    cell = ws.cell(row=ft, column=col0 + j, value=float(s))
                    cell.number_format = FORMATOS[c.fmt][0]
                    cell.font = F_TOTAL
                    cell.fill = RELLENO_TOTAL
            ultima = ft

        for c in cols:
            self.registro.append((hoja or ws.title, nombre_tabla, c))
        return ultima

    # -- hoja completa -------------------------------------------------------------

    def hoja(self, h: Hoja, unidades: str | None = None):
        ws = self.wb.create_sheet(h.nombre)
        self._encabezado(ws, h.titulo, h.descripcion, unidades)
        self.tabla(ws, h.df, h.cols, h.tabla, total_etiqueta=h.total_etiqueta, hoja=h.nombre)
        ws.freeze_panes = ws.cell(row=FILA_HEADER + 1, column=h.congelar_cols + 1)
        return ws

    def diccionario(self):
        filas = [{"Hoja": hj, "Tabla": tb, "Columna": c.header,
                  "Unidad": FORMATOS[c.fmt][1], "Definicion": c.descripcion}
                 for hj, tb, c in self.registro]
        df = pd.DataFrame(filas)
        cols = [Col("Hoja", "Hoja", "texto", "Pestana del libro"),
                Col("Tabla", "Tabla", "texto", "Nombre de la tabla de Excel (leer por este nombre)"),
                Col("Columna", "Columna", "texto", "Encabezado de la columna"),
                Col("Unidad", "Unidad", "texto", "Unidad en que esta guardado el valor"),
                Col("Definicion", "Definicion", "texto_largo", "Que es y de donde sale")]
        ws = self.wb.create_sheet("Diccionario")
        self._encabezado(ws, "Diccionario de datos",
                         "Cada columna de cada tabla del libro: unidad y definicion.",
                         "Se genera de las mismas especificaciones que arman las hojas.")
        self.tabla(ws, df, cols, "tbl_diccionario", hoja="Diccionario")
        ws.freeze_panes = ws.cell(row=FILA_HEADER + 1, column=3)

    def guardar(self, destino: Path):
        destino.parent.mkdir(parents=True, exist_ok=True)
        self.wb.save(destino)
