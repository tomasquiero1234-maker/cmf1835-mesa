"""Pruebas del reporte Excel (reportes/).

Cada control verifica una afirmacion que el reporte hace, contra la fuente:

  1. estructura: 'Datos al' en todas las hojas, una tabla con nombre por hoja,
     encabezado en la fila 6, diccionario completo, tablas legibles por nombre;
  2. conversion a USD: una celda recalculada desde el warehouse;
  3. publicacion vigente: ningun periodo cuenta dos publicaciones;
  4. coherencia entre hojas: el nocional suma lo mismo por instrumento, por
     aseguradora, por contraparte y por subyacente, y los pactos quedan fuera;
  5. entidades legales: una fila por RUT o LEI, sin fusionar identificadores;
  6. conciliacion Confuturo: reproduce EXACTO el resultado validado en 202607;
  7. cuadratura contra el B.8 dentro de +-3% en la industria;
  8. el dashboard no se toco.

    python -m tests.reporte_excel
"""
from __future__ import annotations

import subprocess
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

import openpyxl  # noqa: E402
import pandas as pd  # noqa: E402

from reportes import datos as D  # noqa: E402
from reportes.conciliacion import conciliar  # noqa: E402
from reportes.exportar import generar  # noqa: E402

FALLAS: list[str] = []


def check(ok: bool, texto: str) -> None:
    print(f"  {'OK' if ok else 'XX'} {texto}")
    if not ok:
        FALLAS.append(texto)


def leer_tabla(wb, hoja: str, tabla: str) -> pd.DataFrame:
    """Lee una tabla de Excel POR NOMBRE, como lo haria la automatizacion."""
    ws = wb[hoja]
    ref = ws.tables[tabla].ref
    filas = [[c.value for c in fila] for fila in ws[ref]]
    return pd.DataFrame(filas[1:], columns=filas[0])


def main() -> int:
    salida = Path(tempfile.mkdtemp()) / "reporte.xlsx"
    ruta = generar(salida=salida)
    wb = openpyxl.load_workbook(ruta)
    ctx = D.contexto()
    corte = f"Datos al: {ctx.fecha_cierre:%d-%m-%Y}"

    print("\n[1] estructura")
    check(all(ws["A2"].value == corte for ws in wb.worksheets),
          f"'{corte}' en A2 de las {len(wb.worksheets)} hojas")
    sin_tabla = [ws.title for ws in wb.worksheets if not ws.tables]
    check(not sin_tabla, f"toda hoja tiene al menos una tabla con nombre {sin_tabla or ''}")
    principales = {ws.title: sorted(ws.tables.items(), key=lambda kv: int(''.join(ch for ch in kv[1].split(':')[0] if ch.isdigit())))[0]
                   for ws in wb.worksheets}
    mal_ubicadas = [h for h, (_n, ref) in principales.items() if not ref.split(":")[0].endswith("6")]
    check(not mal_ubicadas, f"la tabla principal de cada hoja empieza en la fila 6 {mal_ubicadas or ''}")
    dic = leer_tabla(wb, "Diccionario", "tbl_diccionario")
    faltan = []
    for ws in wb.worksheets:
        for nombre, ref in ws.tables.items():
            if nombre == "tbl_diccionario":
                continue
            hdr = [c.value for c in ws[ref][0]]
            doc = set(dic[dic.Tabla == nombre].Columna)
            faltan += [f"{nombre}.{h}" for h in hdr if h not in doc]
    check(not faltan, f"el diccionario documenta todas las columnas de todas las tablas {faltan[:5] or ''}")
    legibles = sum(len(ws.tables) for ws in wb.worksheets)
    try:
        for ws in wb.worksheets:
            for nombre in ws.tables:
                leer_tabla(wb, ws.title, nombre)
        check(True, f"las {legibles} tablas se leen por nombre sin error")
    except Exception as e:  # noqa: BLE001
        check(False, f"lectura por nombre: {e}")
    comp = leer_tabla(wb, "Comparativa", "tbl_comparativa")
    check("Dic-2023 (MM USD)" in comp.columns and comp["Dic-2023 (MM USD)"].isna().all(),
          "Dic-2023 existe como columna y esta vacia (no disponible, sin inventar)")

    print("\n[2] conversion a USD")
    stock = leer_tabla(wb, "Stock_Clase", "tbl_stock_clase")
    m = ctx.con.execute(f"""SELECT sum(t.valor_final) FROM raw_bienes_raices t
        JOIN publicacion_vigente v ON v.periodo_informacion=t.periodo_informacion
         AND v.zip_origen=t.zip_origen AND v.recencia=1
        WHERE t.periodo_informacion={ctx.periodo} AND t.rut_compania=96571890""").fetchone()[0]
    esperado = m / ctx.fx[ctx.periodo][0] / 1000
    obs = float(stock.loc[stock["RUT aseguradora"] == 96571890, "Real Estate"].iloc[0])
    check(abs(obs - esperado) < 1e-6,
          f"Real Estate de Confuturo: {obs:,.4f} MM USD = {m:,.0f} M$ / {ctx.fx[ctx.periodo][0]} / 1000")
    check(stock["Sin clasificar"].abs().sum() == 0, "ningun instrumento quedo sin clasificar")

    print("\n[3] publicacion vigente")
    total_vig = sum(float(ctx.con.execute(f"""SELECT COALESCE(sum(t.valor_final),0) FROM {t} t
        JOIN publicacion_vigente v ON v.periodo_informacion=t.periodo_informacion
         AND v.zip_origen=t.zip_origen AND v.recencia=1 WHERE t.periodo_informacion={ctx.periodo}""").fetchone()[0])
        for t in ("raw_equity", "raw_fondo", "raw_extranjero_rf", "raw_extranjero_rv", "raw_otras_inv"))
    total_todas = sum(float(ctx.con.execute(f"SELECT COALESCE(sum(valor_final),0) FROM {t} "
                                            f"WHERE periodo_informacion={ctx.periodo}").fetchone()[0])
                      for t in ("raw_equity", "raw_fondo", "raw_extranjero_rf", "raw_extranjero_rv", "raw_otras_inv"))
    xl = stock[["Equity", "ETF", "Fondos de Inversion", "Fondos Mutuos", "Otros"]].sum().sum()
    rf_ext = D.stock_detalle(ctx, ctx.periodo).query("fuente == 'B.5 renta fija extranjera'").mm_usd.sum()
    check(abs((xl + rf_ext) - total_vig / ctx.fx[ctx.periodo][0] / 1000) < 1e-3,
          f"B.2/B.3/B.5/B.6 del Excel = publicacion vigente ({total_vig / 1e6:,.1f} miles de MM$), "
          f"no la suma de todas ({total_todas / 1e6:,.1f})")

    print("\n[4] coherencia entre hojas de derivados")
    fams = {"CCS": "tbl_ccs", "Swap_Promesa": "tbl_swap_promesa", "Forward_FX": "tbl_forward_fx",
            "Forward_UF": "tbl_forward_uf", "IRS": "tbl_irs", "Opciones": "tbl_opciones", "Futuros": "tbl_futuros"}
    por_inst = sum(pd.to_numeric(leer_tabla(wb, h, t)["Nocional (MM USD)"], errors="coerce").sum()
                   for h, t in fams.items())
    por_aseg = pd.to_numeric(leer_tabla(wb, "Deriv_Aseguradora", "tbl_deriv_aseguradora")
                             ["Nocional total derivados (MM USD)"]).sum()
    por_cp = pd.to_numeric(leer_tabla(wb, "Deriv_Contraparte", "tbl_deriv_contraparte")
                           ["Nocional total derivados (MM USD)"]).sum()
    por_sub = pd.to_numeric(leer_tabla(wb, "Deriv_Subyacente", "tbl_deriv_subyacente")["Nocional (MM USD)"]).sum()
    check(max(por_inst, por_aseg, por_cp, por_sub) - min(por_inst, por_aseg, por_cp, por_sub) < 1e-6,
          f"nocional identico por instrumento, aseguradora, contraparte y subyacente: {por_inst:,.2f} MM USD")
    pactos = pd.to_numeric(leer_tabla(wb, "Pactos", "tbl_pactos")["Nocional (MM USD)"]).sum()
    ops = D.derivados(ctx)
    check(abs(por_aseg - ops[ops.familia != "Pacto"].nocional_mmusd.sum()) < 1e-6 and pactos > 0,
          f"pactos ({pactos:,.2f} MM USD) fuera del total de derivados")

    print("\n[5] entidades legales")
    dc = leer_tabla(wb, "Deriv_Contraparte", "tbl_deriv_contraparte")
    ids = ops[ops.familia != "Pacto"].entidad_id.nunique()
    claves = ops[ops.familia != "Pacto"].contraparte_key.nunique()
    check(len(dc) == ids and dc["ID legal contraparte"].is_unique,
          f"una fila por identificador legal: {len(dc)} entidades (el catalogo del dashboard ve {claves})")
    citi = dc[dc["ID legal contraparte"] == "LEI E57ODZWZ7FF32TWEFA76"]
    check(len(citi) == 1 and citi["Contraparte (nombre legal)"].iloc[0] == "Citibank, National Association",
          "el LEI de Citibank N.A. va a Citibank N.A., no a la agencia en Chile")
    cq = leer_tabla(wb, "Calidad_Contrapartes", "tbl_calidad_contrapartes")
    cred = cq[cq["ID legal contraparte"] == "LEI 54930036B12A3G2SIW61"]
    check(len(cred) == 1 and "CREDIVALORES" in cred.Alerta.iloc[0],
          "el LEI de Credivalores informado como Deutsche Bank queda marcado")

    print("\n[6] conciliacion Confuturo reproduce la validada (202607)")
    r = conciliar(ctx.con, 202607)
    lb = r["libro"]
    esperado_cs = {(pd.Timestamp("2019-08-09"), 4.63): "folio 7215", (pd.Timestamp("2024-06-03"), 5.76): "folio 10314",
                   (pd.Timestamp("2024-11-26"), 4.16): "folio 10666", (pd.Timestamp("2024-11-26"), 4.6): "folio 10665",
                   (pd.Timestamp("2024-10-03"), 2.25): "folio 10552", (pd.Timestamp("2025-02-20"), 5.13): "folio 10814",
                   (pd.Timestamp("2025-03-06"), 5.13): "folio 10838", (pd.Timestamp("2025-03-19"), 5.13): "folio 10851"}
    cs = lb[lb.GROUP == "CS"]
    obtenido = {(row["TRN.DATE"], float(row.RATE)): row.contrapartida_cmf for _, row in cs.iterrows()}
    check(obtenido == esperado_cs, "los 8 swaps calzan con los mismos folios (incluido 10666/10665)")
    rep = lb[lb.GROUP == "REPO"]
    check((rep.resultado == "Cuadra como paquete").all()
          and set(rep.contrapartida_cmf) == {"folio 126", "folio 1140", "folio 129", "folio 1144"},
          "los 13 repos cuadran como 2 paquetes contra los folios 126+1140 y 129+1144")
    check((lb[lb.GROUP.isin(["FXD", "BOND"])].resultado == "Sin reflejo en la CMF").sum() == 7,
          "3 forwards y 4 bonos sin reflejo en la CMF")
    check(set(r["cmf_sin_libro"].folio_operacion) == {"4747", "4748", "4760"},
          "del lado CMF sobran los 3 swaps de 2017 de BBVA Chile (legacy)")
    check(abs(r["tipo_cambio"].iloc[0, 0] - 928.42) < 0.005, "tipo de cambio de Confuturo 928,42, igual al original")
    check(int((lb.resultado != "Sin reflejo en la CMF").sum()) == 21, "21 de 28 filas cuadran")

    print("\n[7] cuadratura contra el B.8")
    tot, b8 = stock["Total"].sum(), stock["Total declarado B.8"].sum()
    check(abs(tot / b8 - 1) < 0.03, f"industria {tot:,.0f} vs declarado {b8:,.0f} MM USD ({(tot / b8 - 1) * 100:+.2f}%)")

    print("\n[8] el dashboard no se toco")
    diff = subprocess.run(["git", "diff", "--quiet", "HEAD", "--", "app/dashboard.py"], cwd=ROOT)
    check(diff.returncode == 0, "app/dashboard.py sin cambios respecto del ultimo commit")

    print("\n" + ("TODO OK" if not FALLAS else f"{len(FALLAS)} FALLAS"))
    for f in FALLAS:
        print("   -", f)
    return 1 if FALLAS else 0


if __name__ == "__main__":
    raise SystemExit(main())
